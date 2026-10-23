from io import BytesIO
from xml.sax.saxutils import escape

from django.contrib.staticfiles import finders
from django.http import HttpResponse

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Image,
    Spacer,
    Table,
    TableStyle,
)

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from .models import Turno


# ==========================================================
# GENERAR PDF
# ==========================================================

def generar_pdf_turno(turno):
    buffer = BytesIO()

    azul = colors.HexColor("#1769AA")
    azul_oscuro = colors.HexColor("#123456")
    gris = colors.HexColor("#64748B")
    fondo = colors.HexColor("#F3F7FB")
    borde = colors.HexColor("#DCE5EF")
    ancho = A4[0] - 4 * cm

    documento = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=1.6 * cm,
        bottomMargin=2 * cm,
    )
    estilos = getSampleStyleSheet()
    estilo_marca = ParagraphStyle(
        "ComprobanteMarca",
        parent=estilos["Title"],
        alignment=TA_LEFT,
        textColor=azul_oscuro,
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=19,
        spaceAfter=3,
    )
    estilo_bajada = ParagraphStyle(
        "ComprobanteBajada",
        parent=estilos["Normal"],
        textColor=gris,
        fontSize=9,
        leading=12,
    )
    estilo_seccion = ParagraphStyle(
        "ComprobanteSeccion",
        parent=estilos["Heading2"],
        textColor=azul_oscuro,
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=15,
        spaceBefore=2,
        spaceAfter=9,
    )
    estilo_etiqueta = ParagraphStyle(
        "ComprobanteEtiqueta",
        parent=estilos["Normal"],
        textColor=gris,
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=11,
    )
    estilo_valor = ParagraphStyle(
        "ComprobanteValor",
        parent=estilos["Normal"],
        textColor=azul_oscuro,
        fontSize=10,
        leading=14,
    )
    estilo_nota = ParagraphStyle(
        "ComprobanteNota",
        parent=estilos["Normal"],
        textColor=gris,
        fontSize=9,
        leading=13,
    )

    usuario = turno.usuario
    cliente = (
        usuario.get_full_name() or usuario.username
        if usuario else "Cliente"
    )
    email = usuario.email if usuario else ""
    nombre_logo = finders.find("images/Imagen molina henze.jpg")
    if nombre_logo is None:
        raise FileNotFoundError(
            "No se encontró el logo del taller en los archivos estáticos."
        )

    logo = Image(nombre_logo, width=2.5 * cm, height=2.5 * cm)
    cabecera = Table(
        [[
            logo,
            [
                Paragraph("TALLER MOLINA HENZE", estilo_marca),
                Paragraph("Mecánica integral", estilo_bajada),
            ],
            Paragraph(
                "<font size='8' color='#64748B'>COMPROBANTE</font>"
                f"<br/><font size='15' color='#123456'><b>"
                f"Turno N.º {escape(str(turno.id))}</b></font>",
                ParagraphStyle(
                    "ComprobanteNumero",
                    parent=estilos["Normal"],
                    alignment=TA_LEFT,
                    leading=21,
                ),
            ),
        ]],
        colWidths=[3 * cm, ancho - 7 * cm, 4 * cm],
    )
    cabecera.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
        ("LINEBELOW", (0, 0), (-1, -1), 2, azul),
    ]))

    estado = Table(
        [[
            Paragraph(
                "<font size='8' color='#DCEBFA'>ESTADO DEL TURNO</font>"
                f"<br/><font size='12' color='#FFFFFF'><b>"
                f"{escape(turno.get_estado_display())}</b></font>",
                ParagraphStyle(
                    "ComprobanteEstado",
                    parent=estilos["Normal"],
                    leading=19,
                ),
            ),
            Paragraph(
                f"<font size='8' color='#DCEBFA'>AVANCE</font>"
                f"<br/><font size='12' color='#FFFFFF'><b>"
                f"{escape(str(turno.porcentaje_reparacion))}%</b></font>",
                ParagraphStyle(
                    "ComprobanteAvance",
                    parent=estilos["Normal"],
                    alignment=TA_LEFT,
                    leading=19,
                ),
            ),
        ]],
        colWidths=[ancho * 0.68, ancho * 0.32],
    )
    estado.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), azul_oscuro),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 14),
        ("RIGHTPADDING", (0, 0), (-1, -1), 14),
        ("TOPPADDING", (0, 0), (-1, -1), 11),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 11),
    ]))

    def campo(etiqueta, valor):
        return [
            Paragraph(escape(etiqueta).upper(), estilo_etiqueta),
            Paragraph(escape(str(valor or "—")), estilo_valor),
        ]

    datos = [
        [campo("Cliente", cliente), campo("Correo electrónico", email)],
        [campo("Vehículo", turno.vehiculo), campo("Patente", turno.patente)],
        [campo("Servicio", str(turno.servicio)), ""],
        [
            campo("Fecha del turno", turno.fecha_turno.strftime("%d/%m/%Y")),
            campo("Horario", turno.fecha_turno.strftime("%H:%M")),
        ],
    ]
    detalle = Table(
        datos,
        colWidths=[ancho / 2, ancho / 2],
    )
    detalle.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), fondo),
        ("GRID", (0, 0), (-1, -1), 0.5, borde),
        ("SPAN", (1, 2), (1, 2)),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 12),
        ("RIGHTPADDING", (0, 0), (-1, -1), 12),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
    ]))

    observaciones = escape(
        turno.observaciones or "Sin observaciones."
    ).replace("\n", "<br/>")
    contenido = [
        cabecera,
        Spacer(1, 16),
        estado,
        Spacer(1, 20),
        Paragraph("Datos del cliente y del vehículo", estilo_seccion),
        detalle,
        Spacer(1, 20),
        Paragraph("Observaciones", estilo_seccion),
        Paragraph(observaciones, estilo_nota),
        Spacer(1, 24),
        Paragraph(
            "Presentá este comprobante al llegar al taller. "
            "La solicitud de turno queda sujeta a confirmación.",
            estilo_nota,
        ),
    ]

    def dibujar_pie(canvas, _documento):
        canvas.saveState()
        canvas.setStrokeColor(borde)
        canvas.setLineWidth(0.7)
        canvas.line(2 * cm, 1.45 * cm, A4[0] - 2 * cm, 1.45 * cm)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(gris)
        canvas.drawString(
            2 * cm,
            1.05 * cm,
            "Taller Molina Henze · Mecánica integral",
        )
        canvas.drawRightString(
            A4[0] - 2 * cm,
            1.05 * cm,
            f"Comprobante de turno N.º {turno.id}",
        )
        canvas.restoreState()

    documento.build(
        contenido,
        onFirstPage=dibujar_pie,
        onLaterPages=dibujar_pie,
    )

    buffer.seek(0)
    response = HttpResponse(buffer, content_type="application/pdf")
    response["Content-Disposition"] = (
        f'attachment; filename="comprobante_turno_{turno.id}.pdf"'
    )
    return response


# ==========================================================
# CREAR EXCEL
# ==========================================================


def crear_excel_turnos(turnos=None):
    from django.db.models import Count
    from openpyxl.styles import (
        Alignment,
        Border,
        Font,
        PatternFill,
        Side,
    )
    from openpyxl.utils import get_column_letter

    buffer = BytesIO()
    libro = Workbook()

    hoja_todos = libro.active
    hoja_todos.title = "Todos los turnos"

    hoja_nuevos = libro.create_sheet("Clientes nuevos")
    hoja_habituales = libro.create_sheet("Clientes habituales")

    encabezados = [
        "ID",
        "Cliente",
        "Email",
        "Teléfono",
        "Vehículo",
        "Patente",
        "Servicio",
        "Fecha",
        "Hora",
        "Tipo de cliente",
        "Estado",
        "Avance",
        "Observaciones",
    ]
    hojas = [hoja_todos, hoja_nuevos, hoja_habituales]
    ancho_columnas = {
        "A": 11,
        "B": 25,
        "C": 30,
        "D": 17,
        "E": 24,
        "F": 15,
        "G": 24,
        "H": 15,
        "I": 12,
        "J": 20,
        "K": 29,
        "L": 12,
        "M": 42,
    }
    azul = "123456"
    azul_medio = "1769AA"
    celeste = "EAF3FB"
    blanco = "FFFFFF"
    gris_borde = "DCE5EF"
    borde_fino = Side(style="thin", color=gris_borde)

    libro.properties.title = "Turnos - Taller Molina Henze"
    libro.properties.subject = "Reporte de turnos y clientes"
    libro.properties.creator = "Taller Molina Henze"

    # Si no recibimos turnos, exportamos todos.
    if turnos is None:
        turnos = Turno.objects.all()

    # Cargar relaciones para evitar consultas repetidas.
    turnos = turnos.select_related(
        "usuario",
        "servicio",
    ).order_by("-fecha_turno")

    # Contar los turnos de cada usuario.
    cantidades = dict(
        Turno.objects.values("usuario_id")
        .annotate(cantidad=Count("id"))
        .values_list("usuario_id", "cantidad")
    )

    # Recuperar teléfonos desde el modelo Cliente, cuando existan.
    from .models import Cliente

    emails = [
        turno.usuario.email
        for turno in turnos
        if turno.usuario and turno.usuario.email
    ]

    clientes_por_email = {
        cliente.email.lower(): cliente
        for cliente in Cliente.objects.filter(email__in=emails)
    }

    for turno in turnos:
        usuario = turno.usuario

        if usuario:
            nombre = usuario.get_full_name() or usuario.username
            email = usuario.email or ""
            cliente = clientes_por_email.get(email.lower())
            telefono = cliente.telefono if cliente else ""
            cantidad = cantidades.get(usuario.id, 0)
        else:
            nombre = "Cliente"
            email = ""
            telefono = ""
            cantidad = 0

        tipo_cliente = (
            "Cliente habitual"
            if cantidad > 1
            else "Cliente nuevo"
        )

        fecha = turno.fecha_turno

        fila = [
            turno.id,
            nombre,
            email,
            telefono,
            turno.vehiculo or "",
            turno.patente or "",
            str(turno.servicio),
            fecha.strftime("%d/%m/%Y"),
            fecha.strftime("%H:%M"),
            tipo_cliente,
            turno.get_estado_display(),
            f"{turno.porcentaje_reparacion or 0}%",
            turno.observaciones or "",
        ]

        hoja_todos.append(fila)

        if tipo_cliente == "Cliente habitual":
            hoja_habituales.append(fila)
        else:
            hoja_nuevos.append(fila)

    titulos = {
        "Todos los turnos": "REPORTE GENERAL DE TURNOS",
        "Clientes nuevos": "TURNOS DE CLIENTES NUEVOS",
        "Clientes habituales": "TURNOS DE CLIENTES HABITUALES",
    }

    for hoja in hojas:
        datos = [
            fila for fila in hoja.values
            if any(valor is not None for valor in fila)
        ]
        hoja.delete_rows(1, hoja.max_row)

        hoja.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=len(encabezados),
        )
        titulo = hoja.cell(row=1, column=1, value=titulos[hoja.title])
        titulo.font = Font(name="Aptos Display", size=16, bold=True, color=blanco)
        titulo.fill = PatternFill("solid", fgColor=azul)
        titulo.alignment = Alignment(vertical="center")
        hoja.row_dimensions[1].height = 34

        hoja.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=len(encabezados),
        )
        descripcion = hoja.cell(
            row=2,
            column=1,
            value=f"Turnos incluidos: {len(datos)}  |  Taller Molina Henze",
        )
        descripcion.font = Font(name="Aptos", size=10, color="52657A")
        descripcion.fill = PatternFill("solid", fgColor=celeste)
        descripcion.alignment = Alignment(vertical="center")
        hoja.row_dimensions[2].height = 23
        hoja.row_dimensions[3].height = 8

        for numero_columna, encabezado in enumerate(encabezados, start=1):
            celda = hoja.cell(row=4, column=numero_columna, value=encabezado)
            celda.font = Font(name="Aptos", size=10, bold=True, color=blanco)
            celda.fill = PatternFill("solid", fgColor=azul_medio)
            celda.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            celda.border = Border(bottom=Side(style="medium", color=azul))
        hoja.row_dimensions[4].height = 30

        for numero_fila, fila in enumerate(datos, start=5):
            for numero_columna, valor in enumerate(fila, start=1):
                celda = hoja.cell(
                    row=numero_fila,
                    column=numero_columna,
                    value=valor,
                )
                celda.font = Font(name="Aptos", size=10, color="263445")
                celda.alignment = Alignment(
                    vertical="top",
                    wrap_text=numero_columna in (2, 3, 5, 7, 13),
                )
                celda.border = Border(bottom=borde_fino)
                if numero_fila % 2 == 0:
                    celda.fill = PatternFill("solid", fgColor="F6F9FC")

            hoja.row_dimensions[numero_fila].height = 30
            estado = hoja.cell(row=numero_fila, column=11)
            estado.fill = PatternFill("solid", fgColor=celeste)
            estado.font = Font(name="Aptos", size=10, bold=True, color=azul)
            hoja.cell(row=numero_fila, column=1).font = Font(
                name="Aptos",
                size=10,
                bold=True,
                color=azul,
            )

        hoja.freeze_panes = "A5"
        hoja.auto_filter.ref = (
            f"A4:{get_column_letter(len(encabezados))}{max(4, hoja.max_row)}"
        )
        hoja.sheet_view.showGridLines = False
        hoja.sheet_view.zoomScale = 85
        hoja.sheet_properties.tabColor = azul_medio
        hoja.print_title_rows = "1:4"
        hoja.page_setup.orientation = "landscape"
        hoja.page_setup.fitToWidth = 1
        hoja.page_setup.fitToHeight = 0
        hoja.sheet_properties.pageSetUpPr.fitToPage = True
        hoja.page_margins.left = 0.25
        hoja.page_margins.right = 0.25
        hoja.page_margins.top = 0.5
        hoja.page_margins.bottom = 0.5

        for letra, ancho in ancho_columnas.items():
            hoja.column_dimensions[letra].width = ancho

    libro.save(buffer)
    buffer.seek(0)

    return buffer

    # ======================================================
    # TURNOS
    # ======================================================

    turnos = Turno.objects.select_related(
        "usuario",
        "servicio"
    ).all()


    for turno in turnos:

        if turno.usuario:

            cliente = (
                turno.usuario.get_full_name()
                or turno.usuario.username
            )

            email = turno.usuario.email

        else:

            cliente = "Cliente"

            email = ""


        cantidad_anterior = Turno.objects.filter(

            usuario=turno.usuario

        ).exclude(

            id=turno.id

        ).count()


        if cantidad_anterior > 0:

            tipo_cliente = "Cliente habitual"

        else:

            tipo_cliente = "Cliente nuevo"


        fila = [

            turno.id,

            cliente,

            email,

            turno.vehiculo,

            turno.patente,

            str(turno.servicio),

            turno.fecha_turno.strftime(
                "%d/%m/%Y"
            ),

            turno.fecha_turno.strftime(
                "%H:%M"
            ),

            tipo_cliente,

            turno.get_estado_display(),

            f"{turno.porcentaje_reparacion}%",

            turno.observaciones,
        ]


        hoja_todos.append(
            fila
        )


        if tipo_cliente == "Cliente nuevo":

            hoja_nuevos.append(
                fila
            )

        else:

            hoja_habituales.append(
                fila
            )


    # ======================================================
    # ANCHO DE COLUMNAS
    # ======================================================

    for hoja in [

        hoja_todos,

        hoja_nuevos,

        hoja_habituales

    ]:

        for columna in hoja.columns:

            maximo = 0

            letra = columna[0].column_letter


            for celda in columna:

                if celda.value:

                    longitud = len(
                        str(celda.value)
                    )

                    if longitud > maximo:

                        maximo = longitud


            hoja.column_dimensions[
                letra
            ].width = min(
                maximo + 2,
                40
            )


    libro.save(
        buffer
    )


    buffer.seek(0)


    return buffer